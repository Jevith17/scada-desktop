import openpyxl, json, os, csv

xlsx_path = '/home/jevith/Downloads/PMS_PROJECT_REAL_DATA- AMCOS PAINTS.xlsx'
wb = openpyxl.load_workbook(xlsx_path, data_only=True)

def to_float(v, default=0.0):
    if v is None:
        return default
    try:
        return float(v)
    except (ValueError, TypeError):
        return default

# 1. Batch Info
ws_batch = wb['1. Batch Info']
batch_rows = list(ws_batch.iter_rows(values_only=True))
start_idx = 0
for idx, r in enumerate(batch_rows):
    if r and r[0] == 'Batch ID':
        start_idx = idx + 1
        break

batches_list = []
for r in batch_rows[start_idx:]:
    if r and r[0]:
        batches_list.append({
            'batch_id': str(r[0]).strip(),
            'product_name': str(r[1]).strip() if r[1] else '',
            'batch_size_kg': to_float(r[2]),
            'start_time': str(r[3]).strip() if r[3] else '',
            'end_time': str(r[4]).strip() if r[4] else '',
            'mixing_duration_min': to_float(r[5]),
            'qc_result': str(r[6]).strip() if r[6] else 'Pass',
            'production_line': str(r[7]).strip() if len(r) > 7 and r[7] else '',
            'paint_type': str(r[8]).strip() if len(r) > 8 and r[8] else '',
            'charging_duration_min': to_float(r[9]) if len(r) > 9 else 0.0,
            'discharge_duration_min': to_float(r[10]) if len(r) > 10 else 0.0,
            'mixing_start_time': str(r[11]).strip() if len(r) > 11 and r[11] else '',
            'mixing_end_time': str(r[12]).strip() if len(r) > 12 and r[12] else ''
        })

print(f'Extracted {len(batches_list)} batches.')

# 2. Process Data
ws_proc = wb['2. Process Data']
proc_rows = list(ws_proc.iter_rows(values_only=True))
proc_start = 0
proc_header = []
for idx, r in enumerate(proc_rows):
    if r and r[0] == 'Timestamp':
        proc_header = [str(c).strip() for c in r if c is not None]
        proc_start = idx + 1
        break

telemetry_by_batch = {b['batch_id']: [] for b in batches_list}

for r in proc_rows[proc_start:]:
    if not r or not r[0] or not r[1]:
        continue
    bid = str(r[1]).strip()
    if bid not in telemetry_by_batch:
        telemetry_by_batch[bid] = []
    
    pt = {
        'timestamp': str(r[0]).strip(),
        'batch_id': bid,
        'phase': str(r[2]).strip() if r[2] else 'MIXING',
        'elapsed_time_min': to_float(r[3]),
        'mixing_time_min': to_float(r[4]),
        'batch_mass_kg': to_float(r[5]),
        'tank_level_pct': to_float(r[6]),
        'inlet_flow_kg_min': to_float(r[7]),
        'outlet_flow_kg_min': to_float(r[8]),
        'temperature_C': to_float(r[9]),
        'agitator_rpm': to_float(r[10]),
        'motor_current_A': to_float(r[11]),
        'motor_power_kW': to_float(r[12]),
        'pressure_bar': to_float(r[13]),
        'viscosity_cP': to_float(r[14]),
        'pH': to_float(r[15], default=None),
        'valve_position_pct': to_float(r[16]),
        'feed_pump_status': str(r[17]).strip() if r[17] else 'STOPPED',
        'agitator_status': str(r[18]).strip() if r[18] else 'STOPPED',
        'temp_sensor_status': str(r[19]).strip() if r[19] else 'NORMAL',
    }
    telemetry_by_batch[bid].append(pt)

total_pts = sum(len(v) for v in telemetry_by_batch.values())
print(f'Extracted {total_pts} telemetry points across {len(telemetry_by_batch)} batches.')

# 3. Fault Library
ws_fault = wb['7. Fault Library']
fault_rows = list(ws_fault.iter_rows(values_only=True))
fault_defs = []
scenarios = []
mode = 'defs'
for r in fault_rows:
    if not r or not any(r):
        continue
    c0 = str(r[0]).strip() if r[0] is not None else ''
    if 'Fault Injection Scenarios' in c0:
        mode = 'scenarios'
        continue
    if 'Confidence Levels' in c0:
        mode = 'confidence'
        continue
    if c0 in ('Fault Library', 'Fault definitions and detection parameters used by the diagnosis engine', 'Fault Code', 'Scenario ID', 'Confidence'):
        continue
    
    if mode == 'defs' and c0:
        fault_defs.append({
            'code': c0,
            'description': str(r[1]).strip() if r[1] else '',
            'affected_equipment': str(r[2]).strip() if r[2] else '',
            'primary_variable': str(r[3]).strip() if r[3] else '',
            'deviation_type': str(r[4]).strip() if r[4] else '',
            'magnitude': str(r[5]).strip() if r[5] else '',
            'ramp_time_min': to_float(r[6]),
            'active_phase': str(r[7]).strip() if r[7] else '',
            'secondary_variables': str(r[8]).strip() if len(r)>8 and r[8] else '',
            'detection_rule': str(r[9]).strip() if len(r)>9 and r[9] else '',
            'severity': str(r[10]).strip() if len(r)>10 and r[10] else 'WARNING',
            'evidence_statement': str(r[11]).strip() if len(r)>11 and r[11] else ''
        })
    elif mode == 'scenarios' and c0 and c0.startswith('SC-'):
        scenarios.append({
            'scenario_id': c0,
            'fault_code': str(r[1]).strip() if r[1] else '',
            'batch_id': str(r[2]).strip() if r[2] else '',
            'injection_time_min': to_float(r[3])
        })

print(f'Extracted {len(fault_defs)} fault definitions, {len(scenarios)} scenarios.')

# 4. Equipment
ws_eq = wb['6. Equipment']
eq_rows = list(ws_eq.iter_rows(values_only=True))
eq_items = []
line_configs = []
mode = 'eq'
for r in eq_rows:
    if not r or not any(r): continue
    c0 = str(r[0]).strip() if r[0] is not None else ''
    if 'Line Configuration' in c0:
        mode = 'lines'
        continue
    if c0 in ('Equipment Details', 'Rated specifications of key equipment', 'Equipment', 'Line'):
        continue
    if mode == 'eq' and c0:
        eq_items.append({
            'equipment': c0,
            'rated_capacity_power': str(r[1]).strip() if r[1] else '',
            'operating_range': str(r[2]).strip() if r[2] else '',
            'key_dimensions': str(r[3]).strip() if len(r)>3 and r[3] else ''
        })
    elif mode == 'lines' and c0 in ('1AP', '2AP', '4AP'):
        line_configs.append({
            'line': c0,
            'paint_type': str(r[1]).strip() if r[1] else '',
            'agitator_power_kW': to_float(r[2]),
            'agitator_speed_rpm': to_float(r[3]),
            'supply_voltage_V': to_float(r[4]),
            'power_factor': to_float(r[5]),
            'rated_current_A': to_float(r[6]),
            'charging_flow_kg_min': to_float(r[7]),
            'discharge_flow_kg_min': to_float(r[8]),
        })

# 5. Setpoints and limits
ws_sp = wb['5. Setpoints & Limits']
sp_rows = list(ws_sp.iter_rows(values_only=True))
setpoints = []
for r in sp_rows:
    if not r or not any(r): continue
    c0 = str(r[0]).strip() if r[0] is not None else ''
    if c0 in ('Setpoints & Operating Ranges', 'Normal operating setpoint and min/max range for each parameter', 'Parameter'):
        continue
    if c0:
        setpoints.append({
            'parameter': c0,
            'unit': str(r[1]).strip() if r[1] else '',
            'setpoint': to_float(r[2]),
            'min': to_float(r[3]),
            'max': to_float(r[4]),
            'alarm_low': to_float(r[5]),
            'alarm_high': to_float(r[6]),
            'applies_to': str(r[7]).strip() if len(r)>7 and r[7] else 'All lines'
        })

# Save JSON and CSV
os.makedirs('data/amcos', exist_ok=True)
with open('data/amcos/amcos_batches.json', 'w') as f:
    json.dump(batches_list, f, indent=2)

with open('data/amcos/amcos_fault_library.json', 'w') as f:
    json.dump({'faults': fault_defs, 'scenarios': scenarios}, f, indent=2)

with open('data/amcos/amcos_telemetry.json', 'w') as f:
    json.dump(telemetry_by_batch, f, indent=2)

with open('data/amcos/amcos_equipment.json', 'w') as f:
    json.dump({'equipment': eq_items, 'lines': line_configs}, f, indent=2)

with open('data/amcos/amcos_setpoints.json', 'w') as f:
    json.dump(setpoints, f, indent=2)

# Write clean CSVs
with open('data/amcos/amcos_all_telemetry.csv', 'w', newline='') as f:
    writer = csv.writer(f)
    writer.writerow(proc_header)
    for b in batches_list:
        for pt in telemetry_by_batch[b['batch_id']]:
            writer.writerow([
                pt['timestamp'], pt['batch_id'], pt['phase'], pt['elapsed_time_min'],
                pt['mixing_time_min'], pt['batch_mass_kg'], pt['tank_level_pct'],
                pt['inlet_flow_kg_min'], pt['outlet_flow_kg_min'], pt['temperature_C'],
                pt['agitator_rpm'], pt['motor_current_A'], pt['motor_power_kW'],
                pt['pressure_bar'], pt['viscosity_cP'], pt['pH'] if pt['pH'] is not None else 'N/A',
                pt['valve_position_pct'], pt['feed_pump_status'], pt['agitator_status'],
                pt['temp_sensor_status']
            ])

# Export clean standalone batch CSV for batch 1AP/0191/26-27
with open('data/amcos/sample_batch_1AP_0191_26_27.csv', 'w', newline='') as f:
    writer = csv.writer(f)
    writer.writerow(proc_header)
    for pt in telemetry_by_batch['1AP/0191/26-27']:
        writer.writerow([
            pt['timestamp'], pt['batch_id'], pt['phase'], pt['elapsed_time_min'],
            pt['mixing_time_min'], pt['batch_mass_kg'], pt['tank_level_pct'],
            pt['inlet_flow_kg_min'], pt['outlet_flow_kg_min'], pt['temperature_C'],
            pt['agitator_rpm'], pt['motor_current_A'], pt['motor_power_kW'],
            pt['pressure_bar'], pt['viscosity_cP'], pt['pH'] if pt['pH'] is not None else 'N/A',
            pt['valve_position_pct'], pt['feed_pump_status'], pt['agitator_status'],
            pt['temp_sensor_status']
        ])

print('Successfully exported AMCOS dataset files and clean sample CSVs!')
